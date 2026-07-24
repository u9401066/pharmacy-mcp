"""Tests for file, SQL, vector, and allowlisted web knowledge connectors."""

import html
import json
import sqlite3
from contextlib import closing
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook

from pharmacy_mcp.application.services.connector_access import ConnectorAccessService
from pharmacy_mcp.config import settings
from pharmacy_mcp.domain.models.provider import ProviderQuery
from pharmacy_mcp.domain.models.response import ResponseStatus
from pharmacy_mcp.infrastructure.api.wcf import WCFClient
from pharmacy_mcp.infrastructure.documents import DocumentStore
from pharmacy_mcp.infrastructure.providers.file import FileKnowledgeProvider
from pharmacy_mcp.infrastructure.providers.registry import (
    ProviderRegistry,
    build_default_registry,
)
from pharmacy_mcp.infrastructure.providers.sql import (
    SQLiteKnowledgeProvider,
    SQLTableMapping,
)
from pharmacy_mcp.infrastructure.providers.vector import (
    VectorKnowledgeProvider,
    VectorSearchClient,
)
from pharmacy_mcp.infrastructure.providers.wcf import WCFKnowledgeProvider
from pharmacy_mcp.infrastructure.providers.web import WebKnowledgeProvider


@pytest.fixture
def document_root(tmp_path: Path) -> Path:
    (tmp_path / "guide.md").write_text(
        "# Anticoagulation\nWarfarin requires INR monitoring.",
        encoding="utf-8",
    )
    (tmp_path / "formulary.csv").write_text(
        "name,rule\nWarfarin,high alert\n",
        encoding="utf-8",
    )

    document = Document()
    document.add_paragraph("Warfarin DOCX guideline")
    document.save(tmp_path / "guideline.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["drug", "stock"])
    sheet.append(["Warfarin", 120])
    workbook.save(tmp_path / "inventory.xlsx")
    return tmp_path


@pytest.mark.asyncio
async def test_file_provider_searches_markdown_csv_docx_and_xlsx(
    document_root: Path,
) -> None:
    provider = FileKnowledgeProvider(
        (document_root,),
        max_bytes=1_000_000,
        max_files=20,
    )

    result = await provider.query(ProviderQuery(text="warfarin", limit=10))

    extensions = {match["extension"] for match in result.data["matches"]}
    assert extensions == {".csv", ".docx", ".md", ".xlsx"}
    assert result.data["files_scanned"] == 4
    assert all(
        match["document_id"].startswith("doc-") for match in result.data["matches"]
    )
    assert all(len(match["text_sha256"]) == 64 for match in result.data["matches"])
    assert all(
        match["char_end"] > match["char_start"] for match in result.data["matches"]
    )

    markdown_match = next(
        match for match in result.data["matches"] if match["extension"] == ".md"
    )
    document = await provider.read_document(markdown_match["document_id"], max_chars=20)
    assert document.content == "# Anticoagulation\nWa"
    assert document.text_sha256 == markdown_match["text_sha256"]
    assert document.char_start == 0
    assert document.char_end == 20
    assert document.truncated is True

    registry = ProviderRegistry()
    registry.register(provider)
    service_result = await ConnectorAccessService(registry).read_document(
        markdown_match["document_id"],
        max_chars=20,
    )
    assert service_result.status is ResponseStatus.OK
    assert service_result.data["content"] == document.content
    assert service_result.sources[0].version == document.text_sha256


@pytest.mark.asyncio
async def test_document_access_service_reports_unknown_id() -> None:
    result = await ConnectorAccessService(ProviderRegistry()).read_document(
        "doc-000000000000000000000000"
    )

    assert result.status is ResponseStatus.ERROR
    assert result.errors[0].code == "file_provider_unavailable"


def test_document_store_rejects_symlink_escape(
    document_root: Path,
    tmp_path: Path,
) -> None:
    outside = tmp_path.parent / "outside-pharmacy-secret.md"
    outside.write_text("warfarin secret", encoding="utf-8")
    link = document_root / "escape.md"
    link.symlink_to(outside)
    store = DocumentStore((document_root,), max_bytes=1000, max_files=10)

    with pytest.raises(ValueError, match="outside an allowed root"):
        store.read(link)

    outside.unlink()


@pytest.mark.asyncio
async def test_sql_provider_uses_read_only_allowlisted_projection(
    tmp_path: Path,
) -> None:
    database = tmp_path / "hospital.sqlite3"
    with closing(sqlite3.connect(database)) as connection:
        connection.execute(
            "CREATE TABLE medications (code TEXT, name TEXT, stock INTEGER, secret TEXT)"
        )
        connection.execute(
            "INSERT INTO medications VALUES (?, ?, ?, ?)",
            ("W001", "Warfarin", 42, "must-not-leak"),
        )
        connection.commit()
    provider = SQLiteKnowledgeProvider(
        database,
        (
            SQLTableMapping(
                table="medications",
                search_columns=("code", "name"),
                output_columns=("code", "name", "stock"),
            ),
        ),
    )

    result = await provider.query(ProviderQuery(text="warfarin"))

    row = result.data["tables"]["medications"][0]
    assert row == {"code": "W001", "name": "Warfarin", "stock": 42}


def test_sql_mapping_rejects_injection_identifier() -> None:
    with pytest.raises(ValueError, match="Unsafe SQL identifier"):
        SQLTableMapping(
            table="medications; DROP TABLE medications",
            search_columns=("name",),
            output_columns=("name",),
        )


@pytest.mark.asyncio
async def test_vector_gateway_receives_only_explicit_vector_filters() -> None:
    async def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"] == "Bearer vector-secret"
        payload = request.content.decode()
        assert "ward-a" in payload
        assert "patient_id" not in payload
        return httpx.Response(
            200,
            json={"results": [{"id": "doc-1", "score": 0.98, "text": "Warfarin"}]},
        )

    provider = VectorKnowledgeProvider(
        VectorSearchClient(
            "https://vector.test/search",
            api_key="vector-secret",
            transport=httpx.MockTransport(handler),
        )
    )
    request = ProviderQuery(
        text="warfarin",
        context={"patient_id": "patient-1", "vector_filters": {"ward": "ward-a"}},
    )

    result = await provider.query(request)

    assert result.data["results"][0]["id"] == "doc-1"


@pytest.mark.asyncio
async def test_web_provider_fetches_only_preconfigured_https_pages() -> None:
    async def handler(request: httpx.Request) -> httpx.Response:
        assert str(request.url) == "https://knowledge.test/warfarin"
        return httpx.Response(
            200,
            headers={"content-type": "text/html", "content-length": "70"},
            text="<html><body><h1>Warfarin</h1><p>Monitor INR.</p></body></html>",
        )

    provider = WebKnowledgeProvider(
        ("https://knowledge.test/warfarin",),
        max_bytes=1000,
        transport=httpx.MockTransport(handler),
    )

    result = await provider.query(ProviderQuery(text="INR"))

    assert result.data["matches"][0]["url"] == "https://knowledge.test/warfarin"
    assert "Monitor INR" in result.data["matches"][0]["snippet"]


def test_web_provider_rejects_non_https_or_credentialed_urls() -> None:
    with pytest.raises(ValueError, match="credential-free HTTPS"):
        WebKnowledgeProvider(("http://127.0.0.1/secret",), max_bytes=1000)
    with pytest.raises(ValueError, match="credential-free HTTPS"):
        WebKnowledgeProvider(("https://user@example.test/secret",), max_bytes=1000)


@pytest.mark.asyncio
async def test_wcf_provider_caches_and_projects_only_allowlisted_fields() -> None:
    calls = 0
    rows = [
        {
            "drug_code": "W001",
            "generic_name": "Warfarin",
            "stock": 42,
            "internal_secret": "must-not-leak",
        },
        {"drug_code": "A001", "generic_name": "Aspirin", "stock": 9},
    ]

    async def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        assert request.headers["soapaction"] == '"urn:test/GetMedicationData"'
        assert b"<GetMedicationData" in request.content
        payload = html.escape(json.dumps(rows))
        return httpx.Response(
            200,
            text=(
                '<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/">'
                "<soap:Body><GetMedicationDataResponse>"
                f"<GetMedicationDataResult>{payload}</GetMedicationDataResult>"
                "</GetMedicationDataResponse></soap:Body></soap:Envelope>"
            ),
        )

    provider = WCFKnowledgeProvider(
        WCFClient(
            "https://wcf.test/service.svc",
            "urn:test/GetMedicationData",
            "GetMedicationData",
            cache_ttl_seconds=300,
            transport=httpx.MockTransport(handler),
        ),
        search_fields=("drug_code", "generic_name"),
        output_fields=("drug_code", "generic_name", "stock"),
    )

    first = await provider.query(ProviderQuery(text="warfarin"))
    second = await provider.query(ProviderQuery(text="W001"))

    assert first.data["matches"] == [
        {"drug_code": "W001", "generic_name": "Warfarin", "stock": 42}
    ]
    assert "internal_secret" not in str(first.data)
    assert first.data["cache_hit"] is False
    assert second.data["cache_hit"] is True
    assert calls == 1


def test_wcf_client_rejects_unsafe_endpoint_or_operation() -> None:
    with pytest.raises(ValueError, match="credential-free HTTPS"):
        WCFClient("http://127.0.0.1/service", "urn:test/action", "GetData")
    with pytest.raises(ValueError, match="safe XML name"):
        WCFClient(
            "https://wcf.test/service",
            "urn:test/action",
            "GetData><Injected",
        )


def test_default_registry_registers_wcf_only_with_complete_contract(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "wcf_service_url", "https://wcf.test/service")
    monkeypatch.setattr(settings, "wcf_soap_action", "urn:test/GetData")
    monkeypatch.setattr(settings, "wcf_operation", "GetData")
    monkeypatch.setattr(settings, "wcf_search_fields", ["name"])
    monkeypatch.setattr(settings, "wcf_output_fields", ["code", "name"])

    providers = {item["id"]: item for item in build_default_registry().catalog()}

    assert providers["wcf"]["registered"] is True
